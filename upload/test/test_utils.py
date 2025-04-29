import os
from typing import List

from openpyxl.cell import Cell
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from accounts.models import PublisherAccount, PSTFUser
from upload.models import UploadFile
from upload.constants import frameworks

from django.test import TestCase

from upload.spreadsheets import ValidateSpreadsheet


def get_cell(value, number_format=None) -> Cell:
    wb = Workbook()
    sheet = Worksheet(wb)
    cell = Cell(sheet, value=value)

    if number_format is not None:
        cell.number_format = number_format

    return cell


class UploadValidatorTestCase(TestCase):
    def setUp(self):
        self.filepath = "upload/test/test.xlsx"
        self.pub1 = PublisherAccount.objects.create(publisher_name='Test publisher')
        PSTFUser.objects.create(name='Test user', publisher=self.pub1)

    def tearDown(self):
        try:
            os.remove(self.filepath)
        except OSError:
            pass

    def getValidator(self, framework: str, data: List):
        upload = UploadFile(framework=framework, publisher=self.pub1)
        upload.save()

        wb = Workbook()
        ws = wb.active

        # keep the first row empty
        ws['A1'] = ''

        # header row
        ws.append([f.name for f in frameworks[upload.framework].fields])

        # data values
        if isinstance(data[0], List):
            for row in data:
                ws.append(row)
        else:
            ws.append(data)

        # Save the file
        wb.save(self.filepath)

        with open(self.filepath, 'rb') as file:
            vds = ValidateSpreadsheet(upload, file)

        return vds
